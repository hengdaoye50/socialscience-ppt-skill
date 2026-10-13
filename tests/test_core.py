"""Run with: python -m unittest discover -s tests -v"""
import json,tempfile,unittest
from pathlib import Path
from pptx import Presentation
from academic_ppt.builder import build,update_charts,graph_layers
from academic_ppt.qa import inspect,audit,slide_xml_hashes

ROOT=Path(__file__).resolve().parents[1]

class EndToEndTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.dir=Path(self.temp.name)
        self.spec=json.loads((ROOT/'examples/demo.json').read_text(encoding='utf8'))
    def tearDown(self):self.temp.cleanup()
    def test_16_layouts_native_objects_and_sources(self):
        o=self.dir/'demo.pptx';build(self.spec,o,spec_dir=ROOT/'examples')
        k=inspect(o)
        self.assertEqual(k['slide_count'],16)
        self.assertEqual(k['chart_count'],1);self.assertEqual(k['table_count'],1)
        self.assertEqual(audit(o,expected=self.spec)['error_count'],0)
        prs=Presentation(o)
        self.assertGreater(sum(1 for sh in prs.slides[6].shapes if sh.shape_type==9),2) # native connectors
    def test_preserves_base_slide_xml(self):
        b=self.dir/'original.pptx';build(self.spec,b,spec_dir=ROOT/'examples')
        target=self.dir/'combined.pptx'
        additions={'theme':'nankai','slides':[self.spec['slides'][6],self.spec['slides'][9]]}
        build(additions,target,base=b,after=3)
        self.assertEqual(len(Presentation(target).slides),18)
        self.assertEqual(slide_xml_hashes(b),slide_xml_hashes(target,16))
        self.assertTrue(any('机制示意' in sh.text for sh in Presentation(target).slides[3].shapes if sh.has_text_frame))
    def test_chart_update(self):
        o=self.dir/'demo.pptx';build(self.spec,o,spec_dir=ROOT/'examples')
        info=next(z for z in inspect(o)['slides'] if z['charts']);id=info['charts'][0]['shape_id']
        edit=[{'slide':info['slide'],'shape_id':id,'categories':['甲','乙','丙','丁','戊','己'],'series':[{'name':'替换一','values':[1,2,3,4,5,6]},{'name':'替换二','values':[7,8,9,10,11,12]}]}]
        dst=self.dir/'updated.pptx';update_charts(o,edit,dst)
        shape=next(s for s in Presentation(dst).slides[info['slide']-1].shapes if s.shape_id==id)
        self.assertEqual(shape.chart.series[0].name,'替换一')
        self.assertEqual(list(shape.chart.series[0].values),[1.,2.,3.,4.,5.,6.])
    def test_cyclic_mechanism_requires_explicit_layers(self):
        with self.assertRaises(ValueError):graph_layers([{'id':'a'},{'id':'b'}],[{'source':'a','target':'b'},{'source':'b','target':'a'}])
    def test_refuse_invalid_estimate(self):
        s={'slides':[{'type':'estimate','title':'bad','items':[{'name':'a','low':3,'estimate':2,'high':4},{'name':'b','low':1,'estimate':2,'high':3}]}]}
        with self.assertRaises(ValueError):build(s,self.dir/'bad.pptx')

if __name__=='__main__':unittest.main()
