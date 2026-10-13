"""Semantic-first PPTX authoring with original, native editable visual elements."""
import json, math, re, os
from pathlib import Path
from collections import defaultdict, deque
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from .theme import Theme

W, H = 13.333, 7.5
MARGIN=0.68
CONTENT_Y=1.46
CONTENT_END=6.78
DEFAULT_FONT = "Noto Sans CJK SC"
ALLOWED = {'cover','section','question','tension','literature','argument','mechanism','process','design','chart','estimate','table','evidence','image','conclusion','closing'}

def col(h): return RGBColor.from_string(h.upper())
def inch(x): return Inches(float(x))
def pt(x): return Pt(float(x))
def req(d,k,t=None):
    if k not in d: raise ValueError(f'Missing required property {k!r} in {d.get("type","slide")}')
    v=d[k]
    if t and not isinstance(v,t): raise ValueError(f'{k} must be {t}')
    return v

def safe_text(v):
    if not isinstance(v, (str,int,float)): raise ValueError('Text must be a string or number')
    return str(v)

def tx(sl, text,x,y,w,h,sz=18,color='242424',bold=False,align='left',font=None,margin=0.0,
       anchor=MSO_ANCHOR.TOP, linespace=1.14):
    t=sl.shapes.add_textbox(inch(x),inch(y),inch(w),inch(h))
    tf=t.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=inch(margin);tf.margin_right=inch(margin)
    tf.margin_top=inch(.02);tf.margin_bottom=inch(.02);tf.vertical_anchor=anchor; tf.auto_size=MSO_AUTO_SIZE.NONE
    for i,line in enumerate(safe_text(text).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_before=pt(0);p.space_after=pt(0);p.line_spacing=1.12
        r=p.add_run();r.text=line;f=r.font;f.name=font or DEFAULT_FONT;f.size=pt(sz);f.bold=bold;f.color.rgb=col(color)
    return t

def rect(sl,x,y,w,h,fill,line=None,radius=False,linewidth=0.8):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, inch(x),inch(y),inch(w),inch(h))
    sh.fill.solid();sh.fill.fore_color.rgb=col(fill)
    if line: sh.line.color.rgb=col(line);sh.line.width=pt(linewidth)
    else: sh.line.fill.background()
    if radius:
        try: sh.adjustments[0]=.11
        except Exception: pass
    return sh

def line(sl,x1,y1,x2,y2,c,width=1.1,dashed=False,arrow=False):
    sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,inch(x1),inch(y1),inch(x2),inch(y2))
    sh.line.color.rgb=col(c);sh.line.width=pt(width)
    if dashed:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    if arrow:
        # Native OOXML arrowhead, editable as PowerPoint connector
        from pptx.oxml.xmlchemy import OxmlElement
        ln=sh.line._get_or_add_ln(); end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm');ln.append(end)
    return sh

def num_circle(sl,x,y,number,t):
    sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,inch(x),inch(y),inch(.38),inch(.38));sh.fill.solid();sh.fill.fore_color.rgb=col(t.primary);sh.line.fill.background()
    tx(sl,number,x,y+.015,.38,.29,12,t.white,True,'center')

def footer(sl,t,n,total,author='',source=''):
    line(sl,MARGIN,7.02,W-MARGIN,7.02,t.line,.7)
    tx(sl,author[:65] or 'ACADEMIC PRESENTATION',MARGIN,7.075,7.6,.23,9.2,t.muted)
    tx(sl,f'{n:02d} / {total:02d}',11.7,7.07,.94,.23,9.2,t.muted,align='right')
    if source:tx(sl,'来源：'+source[:100],MARGIN,6.77,11.82,.20,9.3,t.muted)

def header(sl,t,title,section='',tag=''):
    rect(sl,MARGIN,.38,.08,.47,t.primary)
    tx(sl,title,.9,.37,11.6,.62,28,t.ink,True)
    line(sl,MARGIN,1.17,W-MARGIN,1.17,t.line,.8)
    if section:tx(sl,section,.9,1.22,9.3,.22,10,t.primary,True)
    if tag:tx(sl,tag,10.1,1.23,2.5,.22,10,t.muted,align='right')

def card(sl,t,x,y,w,h,heading,body,index=None,fill=None,heading_sz=17,body_sz=14):
    rect(sl,x,y,w,h,fill or t.soft,radius=True)
    if index is not None:tx(sl,f'{index:02d}',x+.22,y+.17,.45,.32,13,t.primary,True)
    hx=x+.23 if index is None else x+.74
    tx(sl,heading,hx,y+.13,w-(hx-x)-.17,.40,heading_sz,t.deep,True)
    tx(sl,body,x+.23,y+.70,w-.46,h-.82,body_sz,t.body)

def pages_title(title,typ,t,prs):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid();s.background.fill.fore_color.rgb=col(t.white)
    header(s,t,title,typ.upper());return s

def note(sl, data):
    # Notes retain traceability, including full source strings and slide intent.
    try:
        tf=sl.notes_slide.notes_text_frame
        if tf:tf.text=json.dumps({k:v for k,v in data.items() if k not in ('image',)},ensure_ascii=False)
    except Exception:pass

def layout_cover(s,d,t,deck):
    rect(s,0,0,.19,H,t.primary);rect(s,.76,.72,.62,.07,t.primary)
    tx(s,d.get('kicker','学术研究汇报 / RESEARCH PRESENTATION'),.77,1.07,10.6,.36,14,t.primary,True)
    tx(s,req(d,'title'),.77,1.89,11.85,1.45,39,t.ink,True)
    if d.get('subtitle'):tx(s,d['subtitle'],.81,3.66,11.7,.86,20,t.deep)
    line(s,.79,5.52,12.44,5.52,t.line)
    tx(s,d.get('author',deck.get('author','')),.82,5.80,10,.38,17,t.body,True)
    tx(s,d.get('affiliation',deck.get('affiliation','')),.82,6.22,11,.34,13,t.muted)
    tx(s,d.get('date',deck.get('date','')),.82,6.68,6,.28,12,t.muted)
    rect(s,11.44,5.77,1.02,1.02,t.soft,radius=True)
    tx(s,'研 究',11.61,6.10,.72,.29,13,t.primary,True,'center')

def layout_section(s,d,t):
    rect(s,.73,1.49,.09,3.51,t.primary);tx(s,d.get('number','01'),1.03,1.32,2.5,1.28,66,t.primary,True)
    tx(s,req(d,'title'),1.11,3.00,10.9,.9,35,t.ink,True)
    tx(s,d.get('lead',''),1.13,4.15,10.2,.8,18,t.body)

def layout_question(s,d,t):
    header(s,t,d.get('title','研究问题'),'RESEARCH QUESTION')
    tx(s,'经验观察',.9,1.76,5.1,.43,15,t.primary,True)
    tx(s,req(d,'context'),.91,2.30,5.5,2.60,20,t.body)
    rect(s,6.7,1.73,5.78,3.90,t.soft,radius=True)
    tx(s,'需要解释的问题',7.08,2.11,4.90,.39,15,t.primary,True)
    tx(s,req(d,'question'),7.05,2.88,5.0,1.65,25,t.deep,True)
    if d.get('significance'):tx(s,'为何重要：'+d['significance'],.91,5.92,11.5,.58,15,t.body)

def layout_tension(s,d,t):
    header(s,t,d.get('title','经验与理论之间的张力'),'RESEARCH PUZZLE')
    card(s,t,.80,1.75,5.60,3.04,'理论预期',req(d,'expectation'),1,body_sz=18)
    card(s,t,6.88,1.75,5.60,3.04,'经验发现',req(d,'observation'),2,body_sz=18)
    line(s,6.64,3.14,6.66,3.14,t.primary,1.8)
    rect(s,.81,5.21,11.66,1.11,t.deep,radius=True)
    tx(s,req(d,'puzzle'),1.12,5.49,10.96,.57,19,t.white,True)

def layout_literature(s,d,t):
    header(s,t,d.get('title','文献争鸣与解释缺口'),'LITERATURE REVIEW')
    groups=req(d,'groups',list)
    if not (2<=len(groups)<=4):raise ValueError('literature expects 2-4 groups')
    gap=.17; width=(11.70-gap*(len(groups)-1))/len(groups)
    for i,g in enumerate(groups):
        x=.81+i*(width+gap)
        rect(s,x,1.72,width,3.65,t.soft,radius=True)
        rect(s,x,1.72,width,.085,t.primary)
        tx(s,g['heading'],x+.18,2.0,width-.36,.51,17,t.deep,True)
        tx(s,g['claim'],x+.18,2.72,width-.36,1.78,15,t.body)
        tx(s,g.get('authors',''),x+.18,4.72,width-.36,.42,11,t.muted)
    rect(s,.83,5.74,11.67,.65,t.white,line=t.primary,radius=True)
    tx(s,'解释缺口  /  '+req(d,'gap'),1.06,5.91,11.1,.35,15,t.deep,True)

def layout_argument(s,d,t):
    header(s,t,d.get('title','核心论点'),'CORE ARGUMENT')
    rect(s,.81,1.72,11.7,1.27,t.deep,radius=True)
    tx(s,req(d,'claim'),1.11,2.04,11.00,.63,22,t.white,True)
    reasons=req(d,'reasons',list)
    if not 2<=len(reasons)<=4: raise ValueError('argument expects 2-4 reasons')
    gap=.19; width=(11.69-(len(reasons)-1)*gap)/len(reasons)
    for i,z in enumerate(reasons):
        if isinstance(z,str):z={'title':f'论据 {i+1}','body':z}
        card(s,t,.82+i*(width+gap),3.37,width,2.46,z['title'],z['body'],i+1,heading_sz=16,body_sz=14)
    if d.get('implication'):tx(s,'由此可见：'+d['implication'],.92,6.20,11.2,.41,14,t.body)

def graph_layers(nodes,edges):
    ids=[n['id'] for n in nodes]
    if len(set(ids))!=len(ids):raise ValueError('Duplicate mechanism node IDs')
    for e in edges:
        if e['source'] not in ids or e['target'] not in ids:raise ValueError(f'Unresolved mechanism edge {e}')
    if all('layer' in n for n in nodes):return {n['id']:int(n['layer']) for n in nodes}
    degree={i:0 for i in ids};out=defaultdict(list)
    for e in edges:
        degree[e['target']]+=1;out[e['source']].append(e['target'])
    q=deque(i for i in ids if degree[i]==0);levels={i:0 for i in ids};done=0
    while q:
        u=q.popleft();done+=1
        for v in out[u]:
            levels[v]=max(levels[v],levels[u]+1);degree[v]-=1
            if not degree[v]:q.append(v)
    if done<len(ids):
        raise ValueError('Mechanism contains a cycle: specify explicit layer for all nodes or revise causal direction')
    return levels

def layout_mechanism(s,d,t):
    header(s,t,d.get('title','理论机制'),'MECHANISM')
    nodes=req(d,'nodes',list);edges=req(d,'edges',list)
    if not 2<=len(nodes)<=12:raise ValueError('mechanism supports 2-12 nodes; split larger models')
    layers=graph_layers(nodes,edges);counts=defaultdict(int)
    for n in nodes:counts[layers[n['id']]]+=1
    cols=sorted(set(layers.values()));ncol=len(cols)
    if ncol>5 or max(counts.values())>4: raise ValueError('Mechanism requires at most 5 layers / 4 nodes per layer; split diagram')
    colidx={c:i for i,c in enumerate(cols)};col_w=min(2.18,(11.6-.58*(ncol-1))/ncol)
    full_w=ncol*col_w+(ncol-1)*.58;start=(W-full_w)/2
    centers={};positions={}
    for k,c in enumerate(cols):
        these=[n for n in nodes if layers[n['id']]==c]; gap=.25
        box_h=min(.98,(4.25-gap*(len(these)-1))/len(these))
        occupied=len(these)*box_h+gap*(len(these)-1)
        base=1.87+(4.35-occupied)/2
        for j,n in enumerate(these):
            x=start+k*(col_w+.58);y=base+j*(box_h+gap);positions[n['id']]=(x,y,col_w,box_h)
            centers[n['id']]=(x+col_w/2,y+box_h/2)
    # Links first, boxes on top. Edges carry declared semantic relation.
    palette={'causal':(t.primary,False),'temporal':(t.accent,True),'association':(t.muted,True),'organizational':(t.deep,False)}
    for e in edges:
        a=positions[e['source']];b=positions[e['target']];kind=e.get('kind','causal')
        if kind not in palette:raise ValueError('Edge kind must be causal, temporal, association, organizational')
        c,dash=palette[kind]
        if b[0]>a[0]:p=(a[0]+a[2],a[1]+a[3]/2,b[0],b[1]+b[3]/2)
        elif b[0]<a[0]:p=(a[0],a[1]+a[3]/2,b[0]+b[2],b[1]+b[3]/2)
        else:
            i=1 if b[1]>a[1] else -1;p=(a[0]+a[2]/2,a[1]+(a[3] if i>0 else 0),b[0]+b[2]/2,b[1]+(0 if i>0 else b[3]))
        line(s,*p,c,width=1.5,dashed=dash,arrow=True)
        if e.get('label'):
            mx=(p[0]+p[2])/2;my=(p[1]+p[3])/2
            rect(s,mx-.26,my-.15,.52,.26,t.white)
            tx(s,e['label'],mx-.26,my-.14,.52,.21,10,t.muted,align='center')
    for n in nodes:
        x,y,w,h=positions[n['id']];rect(s,x,y,w,h,t.soft,line=t.accent,radius=True)
        tx(s,n['label'],x+.10,y+.16,w-.2,.42,15,t.deep,True,'center')
        if n.get('detail') and h>.70:tx(s,n['detail'],x+.08,y+.57,w-.16,h-.61,10.5,t.body,align='center')
    tx(s,'关系图例：实线＝因果 / 组织关系；虚线＝时间顺序 / 相关关系（具体以连线语义标注为准）',.81,6.41,11.75,.31,10,t.muted)
    if d.get('note'):tx(s,d['note'],.81,6.05,11.7,.29,11,t.body)

def layout_process(s,d,t):
    header(s,t,d.get('title','关键过程追踪'),'PROCESS TRACING')
    steps=req(d,'steps',list)
    if not 2<=len(steps)<=6:raise ValueError('process supports 2-6 steps')
    n=len(steps);width=(11.65-(n-1)*.13)/n
    line(s,1.17,2.56,12.07,2.56,t.accent,1.5)
    for i,st in enumerate(steps):
        x=.82+i*(width+.13);num_circle(s,x+.10,2.37,str(i+1),t)
        tx(s,st.get('time',f'阶段 {i+1}'),x+.06,1.77,width-.12,.42,13,t.primary,True)
        rect(s,x,2.98,width,2.83,t.soft,radius=True)
        tx(s,st['title'],x+.16,3.22,width-.32,.66,16,t.deep,True)
        tx(s,st.get('detail',''),x+.16,4.04,width-.32,1.08,13,t.body)
        if st.get('actor'):tx(s,'行动者：'+st['actor'],x+.16,5.31,width-.32,.33,10.5,t.muted)
    if d.get('takeaway'):tx(s,'过程判断：'+d['takeaway'],.86,6.18,11.5,.50,14,t.body,True)

def layout_design(s,d,t):
    header(s,t,d.get('title','研究设计'),'RESEARCH DESIGN')
    cells=[('研究对象 / 案例',req(d,'case')),('资料与数据',req(d,'materials')),('方法与识别策略',req(d,'strategy')),('研究边界 / 局限',req(d,'limitations'))]
    for i,(head,body) in enumerate(cells):
        x=.81+(i%2)*5.98;y=1.70+(i//2)*2.25
        card(s,t,x,y,5.70,2.06,head,body,i+1,body_sz=15)
    if d.get('question'):tx(s,'围绕的问题：'+d['question'],.87,6.40,11.4,.36,13,t.muted)

def layout_chart(s,d,t):
    header(s,t,d.get('title','研究结果'),'EMPIRICAL RESULT')
    cats=req(d,'categories',list);series=req(d,'series',list)
    if not cats or len(cats)>20:raise ValueError('chart category count 1-20')
    if not 1<=len(series)<=5:raise ValueError('chart supports 1-5 series')
    data=CategoryChartData();data.categories=[safe_text(c) for c in cats]
    for z in series:
        values=req(z,'values',list)
        if len(values)!=len(cats) or not all(isinstance(v,(int,float)) and math.isfinite(v) for v in values):raise ValueError('chart data must be numeric and align categories')
        data.add_series(z.get('name','Series'),values)
    kind=d.get('chart_type','column'); kinds={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE,'area':XL_CHART_TYPE.AREA,'pie':XL_CHART_TYPE.PIE}
    if kind not in kinds:raise ValueError('chart type: column/bar/line/area/pie')
    if kind=='pie' and len(series)!=1:raise ValueError('Pie requires one series')
    chart=s.shapes.add_chart(kinds[kind],inch(1.05),inch(1.62),inch(11.20),inch(4.60),data).chart
    chart.has_legend=len(series)>1;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False
    chart.chart_style=10
    try:
        chart.value_axis.has_major_gridlines=True;chart.value_axis.tick_labels.font.size=pt(12)
        chart.category_axis.tick_labels.font.size=pt(12)
    except Exception:pass
    try:
        palette=[t.primary,t.accent,t.deep,t.muted,'BCA5B7']
        for j,ss in enumerate(chart.series):ss.format.fill.solid();ss.format.fill.fore_color.rgb=col(palette[j%len(palette)]);ss.format.line.color.rgb=col(palette[j%len(palette)])
    except Exception:pass
    if d.get('insight'):tx(s,'主要发现：'+d['insight'],.97,6.27,11.2,.35,13,t.deep,True)

def layout_estimate(s,d,t):
    header(s,t,d.get('title','效应估计与不确定性'),'INTERVAL ESTIMATES')
    items=req(d,'items',list)
    if not 2<=len(items)<=9:raise ValueError('estimate supports 2-9 items')
    for it in items:
        if not (it['low']<=it['estimate']<=it['high']):raise ValueError('low <= estimate <= high required')
    values=[v for it in items for v in (it['low'],it['high'])]+[d.get('reference',0)]
    amin=float(d.get('xmin',min(values)));amax=float(d.get('xmax',max(values)))
    if amax<=amin or any(v<amin or v>amax for v in values):raise ValueError('CI axis range excludes data')
    chart_x=4.12;chart_w=8.03;chart_y=1.98;chart_h=3.98
    scale=lambda v:chart_x+(float(v)-amin)/(amax-amin)*chart_w
    ref=d.get('reference',0);line(s,scale(ref),chart_y-.13,scale(ref),chart_y+chart_h+.18,t.accent,.9,dashed=True)
    for i,it in enumerate(items):
        y=chart_y+(i+.5)*chart_h/len(items)
        tx(s,it['name'],.97,y-.27,2.93,.34,14,t.body,True)
        tx(s,f"{it['estimate']:.2f} [{it['low']:.2f}, {it['high']:.2f}]",.97,y+.07,2.93,.27,10.5,t.muted)
        line(s,scale(it['low']),y,scale(it['high']),y,t.deep,1.5)
        for z in ('low','high'):line(s,scale(it[z]),y-.08,scale(it[z]),y+.08,t.deep,1.1)
        sh=s.shapes.add_shape(MSO_SHAPE.OVAL,inch(scale(it['estimate'])-.055),inch(y-.055),inch(.11),inch(.11));sh.fill.solid();sh.fill.fore_color.rgb=col(t.primary);sh.line.fill.background()
    line(s,chart_x,chart_y+chart_h+.27,chart_x+chart_w,chart_y+chart_h+.27,t.muted,.8)
    for i in range(5):
        v=amin+(amax-amin)*i/4;x=scale(v);line(s,x,chart_y+chart_h+.24,x,chart_y+chart_h+.34,t.muted,.6)
        tx(s,f'{v:g}',x-.3,chart_y+chart_h+.39,.6,.25,10,t.muted,align='center')
    tx(s,d.get('xlabel','估计值与置信区间'),chart_x,6.56,chart_w,.22,10.5,t.muted,align='center')

def layout_table(s,d,t):
    header(s,t,d.get('title','结果汇总'),'EVIDENCE TABLE')
    headers=req(d,'headers',list);rows=req(d,'rows',list);ncol=len(headers);nrow=len(rows)
    if not (2<=ncol<=7 and 1<=nrow<=10):raise ValueError('table max 7 columns and 10 data rows; split bigger tables')
    if any(len(r)!=ncol for r in rows):raise ValueError('Each data row must match headers count')
    w=11.75;h=min(4.9,.48*(nrow+1));x=.79;y=1.72
    table=s.shapes.add_table(nrow+1,ncol,inch(x),inch(y),inch(w),inch(h)).table
    for j in range(ncol):table.columns[j].width=inch(w/ncol)
    table.rows[0].height=inch(.49)
    for i in range(1,nrow+1):table.rows[i].height=inch((h-.49)/nrow)
    highlight=set(d.get('highlight_rows',[]))
    for i,row in enumerate([headers]+rows):
        for j,v in enumerate(row):
            cell=table.cell(i,j);cell.text=safe_text(v);cell.margin_left=inch(.10);cell.margin_right=inch(.08);cell.margin_top=inch(.07);cell.margin_bottom=inch(.04)
            cell.fill.solid();cell.fill.fore_color.rgb=col(t.deep if i==0 else t.soft if i%2==0 or i in highlight else t.white)
            for par in cell.text_frame.paragraphs:
                par.alignment=PP_ALIGN.LEFT
                for rr in par.runs:
                    rr.font.name=t.font_cn;rr.font.size=pt(12 if ncol>5 or nrow>7 else 14);rr.font.bold=(i==0 or i in highlight);rr.font.color.rgb=col(t.white if i==0 else t.deep if i in highlight else t.body)
    if d.get('note'):tx(s,d['note'],.89,min(6.31,y+h+.21),11.5,.33,12,t.body)

def layout_evidence(s,d,t):
    header(s,t,d.get('title','关键证据'),'PRIMARY EVIDENCE')
    rect(s,.81,1.75,11.71,2.88,t.soft,radius=True)
    tx(s,'“',1.04,1.79,.57,.65,41,t.primary,True)
    tx(s,req(d,'excerpt'),1.57,2.23,10.27,1.71,21,t.ink)
    tx(s,'材料来源：'+req(d,'source'),1.10,4.18,10.7,.25,11,t.muted)
    tx(s,'证据指向',.95,5.03,2.1,.34,15,t.primary,True)
    tx(s,req(d,'interpretation'),.96,5.55,11.26,.76,17,t.body,True)

def layout_image(s,d,t,base_dir):
    header(s,t,d.get('title','研究图片 / 地图'),'DOCUMENTARY MATERIAL')
    p=Path(req(d,'image')).expanduser()
    if not p.is_absolute():p=(base_dir/p).resolve()
    if not p.exists():raise FileNotFoundError(f'Image not found: {p}')
    from PIL import Image
    with Image.open(p) as im: iw,ih=im.size
    maxw,maxh=10.0,4.36;f=min(maxw/iw,maxh/ih);ww=iw*f;hh=ih*f
    # Use inch-converted aspect, not raw pixels.
    ratio=iw/ih;ww=min(maxw,maxh*ratio);hh=ww/ratio
    s.shapes.add_picture(str(p),inch((W-ww)/2),inch(1.55+(4.55-hh)/2),width=inch(ww),height=inch(hh))
    tx(s,req(d,'caption'),.87,6.22,11.6,.35,12,t.body)
    if not d.get('source'):raise ValueError('Research image must include source, copyright/permission or explicitly synthetic status')

def layout_conclusion(s,d,t):
    header(s,t,d.get('title','研究结论与贡献'),'CONCLUSION')
    findings=req(d,'findings',list)
    if not 2<=len(findings)<=4:raise ValueError('conclusion expects 2-4 findings')
    h=3.42/len(findings)
    for i,f in enumerate(findings):
        yy=1.65+i*(h+.09);num_circle(s,.96,yy+.16,str(i+1),t)
        tx(s,f,1.58,yy+.11,10.62,h-.13,18,t.ink,True)
        if i<len(findings)-1:line(s,1.55,yy+h+.01,12.36,yy+h+.01,t.line,.8)
    rect(s,.86,5.40,7.22,1.05,t.soft,radius=True);rect(s,8.30,5.40,4.16,1.05,t.soft,radius=True)
    tx(s,'理论贡献',1.06,5.55,1.3,.32,14,t.primary,True)
    tx(s,d.get('contribution',''),2.31,5.53,5.45,.74,13,t.body)
    tx(s,'研究局限',8.49,5.55,1.28,.31,14,t.primary,True)
    tx(s,d.get('limitations',''),9.77,5.52,2.42,.74,11.5,t.body)

def layout_closing(s,d,t):
    rect(s,0,0,.17,H,t.primary);rect(s,.80,1.41,.75,.07,t.primary)
    tx(s,req(d,'takeaway'),.80,2.01,11.65,1.78,34,t.ink,True)
    tx(s,d.get('thanks','谢谢，欢迎批评指正。'),.82,4.97,10.70,.58,21,t.primary,True)
    line(s,.82,6.28,12.35,6.28,t.line)
    tx(s,d.get('contact',''),.83,6.54,10.84,.37,13,t.muted)

LAYOUTS={'cover':layout_cover,'section':layout_section,'question':layout_question,'tension':layout_tension,'literature':layout_literature,'argument':layout_argument,'mechanism':layout_mechanism,'process':layout_process,'design':layout_design,'chart':layout_chart,'estimate':layout_estimate,'table':layout_table,'evidence':layout_evidence,'image':layout_image,'conclusion':layout_conclusion,'closing':layout_closing}

def build(spec, output, base=None, after=None, spec_dir=None):
    if isinstance(spec,(str,Path)):
        p=Path(spec).resolve();spec_dir=p.parent;spec=json.loads(p.read_text(encoding='utf8'))
    spec_dir=Path(spec_dir or '.').resolve()
    deck=spec.get('deck',{});theme=Theme.from_spec(spec.get('theme','nankai'))
    global DEFAULT_FONT
    DEFAULT_FONT=theme.font_cn
    if not isinstance(spec.get('slides'),list) or not spec['slides']:raise ValueError('slides must be a non-empty list')
    if base:
        prs=Presentation(str(base)); ratio=prs.slide_width/prs.slide_height
        if abs(ratio-16/9)>.03:raise ValueError('Existing PowerPoint must be 16:9; do not distort user slides')
        if abs(prs.slide_width/inch(1)-W)>.1:raise ValueError('Existing slide canvas must be standard widescreen 13.333 x 7.5 inches')
        pre_count=len(prs.slides)
    else:
        prs=Presentation();prs.slide_width=inch(W);prs.slide_height=inch(H);pre_count=0
    count=len(spec['slides']);new=[]
    for n,d in enumerate(spec['slides'],1):
        typ=req(d,'type');
        if typ not in ALLOWED:raise ValueError(f'Unknown slide type {typ}')
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=col(theme.white)
        if typ=='cover':layout_cover(s,d,theme,deck)
        elif typ=='section':layout_section(s,d,theme)
        elif typ=='image':layout_image(s,d,theme,spec_dir)
        else:LAYOUTS[typ](s,d,theme)
        if typ not in ('cover','closing','section'):
            footer(s,theme,n,count,deck.get('affiliation',''),d.get('source',''))
        note(s,d);new.append(s)
    if base and after is not None:
        if after<0 or after>pre_count:raise ValueError('after must be between 0 and original slide count')
        ids=prs.slides._sldIdLst
        # The new slide elements occupy the tail; move only their relationship IDs.
        last=[ids[-count+i] for i in range(count)]
        for elem in last:ids.remove(elem)
        for i,elem in enumerate(last):ids.insert(after+i,elem)
    out=Path(output);out.parent.mkdir(parents=True,exist_ok=True);prs.save(str(out))
    return {'file':str(out),'created_slides':count,'total_slides':len(prs.slides),'base_slides':pre_count}

def update_charts(src,edits,dst):
    """Update native chart data by exact slide (1-based) and shape_id without changing design."""
    prs=Presentation(str(src));edits=json.loads(Path(edits).read_text(encoding='utf8')) if not isinstance(edits,list) else edits
    for e in edits:
        idx=int(e['slide'])-1
        if idx<0 or idx>=len(prs.slides):raise ValueError('Invalid slide number')
        sh=next((s for s in prs.slides[idx].shapes if s.shape_id==int(e['shape_id'])),None)
        if sh is None or not sh.has_chart:raise ValueError(f'No native chart at slide {idx+1}, shape {e["shape_id"]}')
        categories=e['categories'];series=e['series']
        if len(series)!=len(sh.chart.series):raise ValueError('Series count cannot change in template-safe update')
        cd=CategoryChartData();cd.categories=categories
        for ss in series:
            if len(ss['values'])!=len(categories):raise ValueError('Series length mismatch')
            cd.add_series(ss['name'],ss['values'])
        sh.chart.replace_data(cd)
    Path(dst).parent.mkdir(parents=True,exist_ok=True);prs.save(str(dst))
    return len(edits)
